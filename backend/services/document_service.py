import pdfplumber
import docx
from html import escape
from typing import Tuple

try:
    import mammoth
except ImportError:
    mammoth = None

def _extract_text_from_pdf(file_path: str) -> Tuple[int, int, str]:
    """Extracts text from a PDF file."""
    text = ""
    page_count = 0
    with pdfplumber.open(file_path) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return page_count, len(text), text

def _extract_text_from_docx(file_path: str) -> Tuple[int, int, str]:
    """Extracts text from a DOCX file."""
    doc = docx.Document(file_path)
    parts = [para.text for para in doc.paragraphs if para.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append("\t".join(cells))

    text = "\n".join(parts)
    # DOCX files don't have a concept of pages in the same way PDFs do
    return 0, len(text), text

def extract_html_from_docx(file_path: str) -> tuple[str, list[str]]:
    """Converts a DOCX file to semantic HTML for formatted previews."""
    if mammoth is not None:
        with open(file_path, "rb") as docx_file:
            result = mammoth.convert_to_html(docx_file)

        warnings = [message.message for message in result.messages]
        return result.value, warnings

    return _extract_docx_html_fallback(file_path), [
        "Using basic DOCX HTML preview because mammoth is not installed."
    ]

def _run_to_html(run) -> str:
    text = escape(run.text or "")
    if not text:
        return ""

    if run.bold:
        text = f"<strong>{text}</strong>"
    if run.italic:
        text = f"<em>{text}</em>"
    if run.underline:
        text = f"<u>{text}</u>"

    return text

def _paragraph_to_html(paragraph) -> str:
    content = "".join(_run_to_html(run) for run in paragraph.runs).strip()
    if not content:
        return ""

    style_name = (paragraph.style.name if paragraph.style else "").lower()
    if "heading 1" in style_name or style_name == "title":
        return f"<h1>{content}</h1>"
    if "heading 2" in style_name:
        return f"<h2>{content}</h2>"
    if "heading 3" in style_name:
        return f"<h3>{content}</h3>"
    if "list" in style_name:
        return f"<ul><li>{content}</li></ul>"

    return f"<p>{content}</p>"

def _table_to_html(table) -> str:
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            cell_content = "<br />".join(
                escape(paragraph.text.strip())
                for paragraph in cell.paragraphs
                if paragraph.text.strip()
            )
            cells.append(f"<td>{cell_content}</td>")
        rows.append(f"<tr>{''.join(cells)}</tr>")

    return f"<table><tbody>{''.join(rows)}</tbody></table>" if rows else ""

def _extract_docx_html_fallback(file_path: str) -> str:
    document = docx.Document(file_path)
    parts = [_paragraph_to_html(paragraph) for paragraph in document.paragraphs]
    parts.extend(_table_to_html(table) for table in document.tables)

    html = "\n".join(part for part in parts if part)
    return html or "<p>No formatted preview could be generated for this document.</p>"

def _extract_text_from_txt(file_path: str) -> Tuple[int, int, str]:
    """Extracts text from a TXT file."""
    encodings = ["utf-8-sig", "utf-8", "utf-16", "utf-16-le", "utf-16-be", "cp1252", "latin-1"]
    last_error = None

    for encoding in encodings:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                text = f.read()
            return 0, len(text), text
        except UnicodeDecodeError as error:
            last_error = error

    raise ValueError(f"Could not decode text file: {last_error}")

def extract_text_from_file(file_path: str, mime_type: str) -> Tuple[int, int, str]:
    """
    Extracts text from a file based on its MIME type.
    
    Returns a tuple of (page_count, char_count, text).
    """
    if mime_type == "application/pdf":
        return _extract_text_from_pdf(file_path)
    elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _extract_text_from_docx(file_path)
    elif mime_type == "text/plain":
        return _extract_text_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file type: {mime_type}")

